# work_plan_loader.py
"""
Work Plan Data Loader Module — mirrors bsc_loader.py's shape.

Produces ONE document per (Division Objective, Unit) spanning the FULL
FISCAL YEAR (Q1-Q4 together), not one document per quarter. An employee's
annual PMS goal should reflect the whole year's trajectory — e.g. an
objective that ramps from 5% in Q1 to 60% by Q4 is one yearly objective
with a growth curve, not four disconnected quarterly fragments.

No Perspective / Strategic Objective lineage tagging — dropped per
explicit instruction: it required cross-referencing a separate LOS
cascade file that often had no match for a given Division Objective,
producing "N/A" clutter on every document. This loader only reports
what actually comes from the Work Plan docx itself: Department, Fiscal
Year, Unit, Division Objective, and the full-year monthly targets.
"""
import re
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import docx
import logging

try:                                    # package import (API) or flat import (scripts)
    from common import registry
except ImportError:                     # pragma: no cover
    from ..common import registry       # type: ignore

logging.basicConfig(level=logging.INFO, format='%(message)s')
logger = logging.getLogger(__name__)

_QUARTER_ORDER = ("Q1", "Q2", "Q3", "Q4")


def _plan_name(text: str) -> str:
    """'Mobile Banking Business Work plan' -> 'Mobile Banking Business'."""
    return re.sub(r"\s*Work\s*plan.*$", "", text or "", flags=re.IGNORECASE).strip()


def resolve_plan_owner(name: str) -> Tuple[str, str, str, str]:
    """
    Who a plan belongs to, from its title: (level, division, department, unit).
    level is 'division' | 'department' | 'unit' | 'unknown'. Exact names are tried first, so
    'Merchant and Agent Management' is the DEPARTMENT, not the unit 'Agent Management' that
    happens to be written inside it.
    """
    name = (name or "").strip()
    if not name:
        return "unknown", "", "", ""
    div = registry.canon_division(name)
    if div:
        return "division", div, "", ""
    dept = registry.canon_department(name)
    if dept:
        return "department", registry.division_of_department(dept) or "", dept, ""
    owner = registry.department_of_unit(name)               # exact unit name or alias
    if owner:
        hit = registry.find_unit_anywhere(name, owner[0])
        if hit and hit[1] == owner[1]:
            return "unit", hit[0], hit[1], hit[2]
    # the title wraps a name in extra words: look inside it
    div = registry.find_division(name)
    if div and not registry.find_department(name, div):
        return "division", div, "", ""
    dept = registry.find_department(name)
    if dept:
        return "department", registry.division_of_department(dept) or "", dept, ""
    hit = registry.find_unit_anywhere(name)
    if hit:
        return "unit", hit[0], hit[1], hit[2]
    return "unknown", "", name, ""


def _clean_target_phrase(text: str) -> str:
    """
    Strip a trailing mechanism clause ("...by strengthening relationship")
    from a Work Plan monthly target line, keeping the outcome + number
    intact (e.g. "Increase active merchants by 16.83% from the base line").
    """
    match = re.search(r"\bby\s+\w+ing\b", text)
    if match:
        return text[: match.start()].strip().rstrip(",")
    return text.strip()


class WorkPlanLoader:
    """Load and process a Work Plan docx into embedding-ready YEARLY documents."""

    def __init__(self, file_path: Path, config: Dict = None):
        """
        Args:
            file_path: Path to the Work Plan .docx file
            config: Configuration dict (quarter labels — see config.py)
        """
        self.file_path = Path(file_path)
        self.config = config or {}
        self.doc = None
        self.owner_level = "unknown"       # division | department | unit | unknown (from the title)
        self.division = ""
        self.department = None
        self.unit = ""
        self.fiscal_year = None
        self.major_objectives: List[str] = []
        self.entries: List[dict] = []  # flat list, each tagged with unit + quarter

    # =========================================================
    # LOAD
    # =========================================================

    def load(self) -> "WorkPlanLoader":
        logger.info(f"\n📂 Loading file: {self.file_path.name}")
        try:
            self.doc = docx.Document(self.file_path)
            logger.info(f"  ✓ Loaded {len(self.doc.tables)} tables")
        except Exception as e:
            logger.error(f"❌ Error loading workbook: {e}")
        return self

    # =========================================================
    # CLEAN / PARSE
    # =========================================================

    def _parse_table0(self, table) -> dict:
        """
        Table 0 alternates a header-label row with a content row:
          row 0: "Brief description..." (label)   row 1: description content
          row 2: "Major Objectives..." (label)     row 3: objectives content
          row 4: budget column labels               row 5: budget figures
        """
        description_lines, major_objectives = [], []
        budget = {"operational": "", "capital": "", "approved_project": ""}

        if len(table.rows) > 1:
            description_lines = [l.strip() for l in table.rows[1].cells[0].text.split("\n") if l.strip()]
        if len(table.rows) > 3:
            major_objectives = [
                l.strip().rstrip(";") for l in table.rows[3].cells[0].text.split("\n") if l.strip()
            ]
        if len(table.rows) > 5:
            figures = [c.text.strip() for c in table.rows[5].cells]
            if len(figures) >= 3:
                budget = {"operational": figures[0], "capital": figures[1], "approved_project": figures[2]}

        return {
            "description": description_lines,
            "major_objectives": [o for o in major_objectives if o],
            "budget": budget,
        }

    def _parse_quarter_table(self, table) -> List[dict]:
        """
        Row 0: ["Division Objectives", quarter_label, quarter_label, quarter_label, ""]
        Row 1: ["Division Objectives", month1, month2, month3, "Unit"]   <- Unit column is OPTIONAL
        Row 2+: [division_objective, target1, target2, target3, unit?]

        A plan may have NO Unit column (4 columns). Its rows then belong to whoever the TITLE
        names: a unit-level plan gives every row that unit; a department or division plan
        gives every row a blank unit (department-wide / division-wide).
        """
        quarter_labels = self.config.get("QUARTER_LABELS", ("1st Q", "2nd Q", "3rd Q", "4th Q"))
        quarter_numbers = self.config.get(
            "QUARTER_NUMBERS", {"1st Q": "Q1", "2nd Q": "Q2", "3rd Q": "Q3", "4th Q": "Q4"}
        )

        header_row0 = [c.text.strip() for c in table.rows[0].cells]
        header_row1 = [c.text.strip() for c in table.rows[1].cells]
        quarter_label = next((h for h in header_row0 if h in quarter_labels), "")
        quarter = quarter_numbers.get(quarter_label, quarter_label or "Unknown")
        months = header_row1[1:4]
        has_unit_col = len(header_row1) >= 5 and header_row1[4].strip().lower() == "unit"

        entries = []
        for row in table.rows[2:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 4 or not cells[0]:
                continue
            division_objective, m1, m2, m3 = cells[0], cells[1], cells[2], cells[3]
            unit_raw = (cells[4] if has_unit_col and len(cells) > 4 else "").strip()
            division, department = self._owner_of_row(unit_raw)
            unit = self._unit_of_row(unit_raw, division, department)
            entries.append({
                "division_objective": division_objective,
                "quarter": quarter,
                "months": dict(zip(months, [m1, m2, m3])),
                "unit": unit,
                "department": department,
                "division": division,
            })
        return entries

    def _unit_of_row(self, unit_raw: str, division: str, department: str) -> str:
        """The row's unit. No Unit cell: the plan's own unit (blank for a department/division
        plan). A "Both ..." cell becomes 'Both' - finalize() later splits it into one document
        per unit. Any other spelling is mapped to the registry's unit name so that the same
        unit written two ways ('Agent' / 'agent') is one unit, not two."""
        if not unit_raw:
            return self.unit
        key = registry.loose_key(unit_raw)
        if key == registry.MULTI_UNIT_PREFIX or key.startswith(registry.MULTI_UNIT_PREFIX + " "):
            return "Both"
        return registry.canon_unit(unit_raw, department, division or None) or unit_raw

    def _owner_of_row(self, unit_raw: str) -> Tuple[str, str]:
        """(division, department) a row is filed under. A unit written in the row wins, when
        the registry knows it ('Agent' -> Merchant and Agent Management); otherwise the
        plan's own owner, taken from its title."""
        if unit_raw:
            owner = registry.department_of_unit(unit_raw)       # (division, department) or None
            if owner:
                return owner
        return self.division, self.department or ""

    def clean(self) -> "WorkPlanLoader":
        """Parse the loaded docx into department metadata + flat unit-tagged entries."""
        if self.doc is None:
            logger.error("❌ Work Plan not loaded yet")
            return self

        logger.info("\n📋 Cleaning Work Plan data...")

        title_paragraphs = [p.text.strip() for p in self.doc.paragraphs if p.text.strip()]
        dept_line = next(
            (t for t in title_paragraphs
             if re.search(r"work\s*plan", t, re.IGNORECASE)
             and not re.match(r"work\s*plan\s+for\b", t, re.IGNORECASE)),
            "",
        )
        title_name = _plan_name(dept_line)
        if not title_name:  # fallback: filename "MB_Work_Plan_-2026-27FY.docx" -> "MB"
            title_name = re.split(
                r"_?work[_\s-]*plan", self.file_path.stem, flags=re.IGNORECASE
            )[0].strip("_ -")
        self.owner_level, self.division, self.department, self.unit = resolve_plan_owner(title_name)
        if self.owner_level == "unknown":
            logger.warning(f"  ⚠️ '{title_name}' is not in the organisation registry")
        fy_line = next((t for t in title_paragraphs if re.search(r"\d{4}/\d{2,4}", t)), "")
        fy_match = re.search(r"\d{4}/\d{2,4}", fy_line)
        self.fiscal_year = fy_match.group(0) if fy_match else ""

        tables = self.doc.tables
        if tables:
            table0 = self._parse_table0(tables[0])
            self.major_objectives = table0["major_objectives"]

        self.entries = []
        for t in tables[1:]:
            self.entries.extend(self._parse_quarter_table(t))

        # A plan must belong to someone: a division, a department or a unit. Rows that end up
        # with no owner at all can never match a query. (A name the registry does not know is
        # NOT an error here: it is kept as written and finalize() rejects it with a clear message.)
        if self.owner_level != "division":
            ownerless = [e for e in self.entries if not e["department"]]
            if ownerless or (not self.department and not self.entries):
                raise ValueError(
                    f"Could not determine who {self.file_path.name!r} belongs to: its title and file "
                    f"name name no division, department or unit. Put the department or unit in the "
                    f"title (e.g. 'Mobile Banking Work plan') or add the name to "
                    f"Data_Ingestion/common/org_registry.json."
                )

        logger.info(f"  ✅ Cleaned {len(self.entries)} Work Plan entries "
                    f"for {self.department or self.division} ({self.fiscal_year})")
        return self

    # =========================================================
    # DOCUMENTS — ONE per (Division Objective, Unit), spanning the
    # FULL FISCAL YEAR (Q1-Q4 rolled up together)
    # =========================================================

    def _group_into_yearly_objectives(self) -> Dict[Tuple[str, str, str, str], Dict[str, dict]]:
        """
        Group flat quarterly entries by (division_objective, unit, department, division), merging
        all 4 quarters' monthly targets under each key — this is what turns
        4 separate quarter-scoped rows into 1 yearly objective.
        """
        grouped: Dict[Tuple[str, str, str, str], Dict[str, dict]] = {}
        for entry in self.entries:
            key = (entry["division_objective"], entry["unit"], entry["department"], entry["division"])
            grouped.setdefault(key, {})[entry["quarter"]] = entry["months"]
        return grouped

    def to_documents(self) -> list:
        """
        Convert cleaned Work Plan entries into ONE yearly document per
        (Division Objective, Unit) pair, listing every distinct activity
        for the year as a flat list — chronological order preserved, but
        month/quarter labels dropped since this is a yearly objective, not
        a month-by-month schedule.

        Exact-duplicate activity lines (the same recurring task appearing
        in multiple months, e.g. a quarterly check-in task) collapse into
        one bullet so the list doesn't repeat itself. Lines that differ —
        including a numeric progression like "...by 5%" -> "...by 10%" ->
        "...by 15%" — are NOT duplicates and are all kept, in order, so
        the full growth trajectory across the year is still visible.
        """
        if not self.entries:
            self.clean()

        grouped = self._group_into_yearly_objectives()

        documents = []
        for (division_objective, unit, department, division), quarters in grouped.items():
            seen: set[str] = set()
            activities: List[str] = []
            for q in _QUARTER_ORDER:
                months = quarters.get(q)
                if not months:
                    continue
                for month_name in months:  # dict preserves insertion order (chronological)
                    cleaned = _clean_target_phrase(re.sub(r"\s*\n\s*", " ", months[month_name]))
                    if cleaned and cleaned not in seen:
                        seen.add(cleaned)
                        activities.append(cleaned)

            activities_text = "\n".join(f"  - {a}" for a in activities)

            # The text is only what the model needs: the objective + its activities.
            # Department / fiscal year / unit are already in the employee's query, and
            # stay in `metadata`, which is what retrieval filters on.
            text = (
                f"Division Objective (Full Year): {division_objective}\n"
                f"Activities for the Year:\n{activities_text}"
            )
            metadata = {
                "source": "WorkPlan",
                "division": division,
                "department": department,
                "fiscal_year": self.fiscal_year,
                "unit": unit,
                "division_objective": division_objective,
            }
            documents.append({"text": text.strip(), "metadata": metadata})
        return documents