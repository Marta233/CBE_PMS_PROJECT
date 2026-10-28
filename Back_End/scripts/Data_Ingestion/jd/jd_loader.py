"""
jd_loader.py - ONE Job Description loader for every division.

The file FORMAT decides how it is read (a Word file holds each JD in a table, a PDF holds them
as text blocks). The DIVISION never decides anything: division, department and unit are read
from each JD's own fields and mapped to canonical names by the organisation registry when the
document is ingested. A new division's JDs therefore need no new loader.
"""
import re
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd

try:                                    # package import (API) or flat import (scripts)
    from common import registry
except ImportError:                     # pragma: no cover
    from ..common import registry       # type: ignore

from .jd_config import (
    INFER_DEPARTMENT_FROM_UNIT,
    RESP_START_KEYWORDS,
    RESP_STOP_KEYWORDS,
)

# ─────────────────────────────────────────────────────────────
# Shared helpers
# ─────────────────────────────────────────────────────────────

_MOJIBAKE = {
    "â€™": "'", "â€œ": '"', "â€\x9d": '"', "â€“": "-", "â€”": "-",
    "Â": "", "ﬁ": "fi", "ﬂ": "fl",
}


def _clean_text(text: str) -> str:
    """Normalize extracted Word/PDF text."""
    for bad, good in _MOJIBAKE.items():
        text = text.replace(bad, good)
    text = re.sub(r"[•●■]", "", text)
    text = re.sub(r"[^\w\s,.;:/&()'\-]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _parse_grade(raw: Optional[str]) -> Optional[int]:
    if not raw:
        return None
    m = re.search(r"(\d+)", raw)
    return int(m.group(1)) if m else None


def _is_stop(text: str) -> bool:
    lo = text.lower()
    return any(kw.lower() in lo for kw in RESP_STOP_KEYWORDS)


def _is_resp_header(text: str) -> bool:
    lo = text.lower()
    return any(kw.lower() in lo for kw in RESP_START_KEYWORDS)


def _infer_department(unit: str, reports_to: str, division: str) -> str:
    """Department from the registry, via the unit (or the boss) the JD names. '' when unsure."""
    div = division or None
    for text in (unit, reports_to):
        hit = registry.find_unit_anywhere(text, div)
        if hit:
            return hit[1]
    return registry.find_department(reports_to, div) or ""


# ─────────────────────────────────────────────────────────────
# WORD (.docx) reader - one JD per table
# ─────────────────────────────────────────────────────────────

def _parse_docx_table(table) -> Optional[Dict]:
    field_paras: Dict[str, str] = {}
    resp_cell = None

    # A merged cell appears once per grid column it spans. De-duplicate on the underlying XML
    # element (kept alive in `seen_cells`). id(cell) is NOT safe: python-docx builds a new
    # temporary wrapper each time and Python reuses ids of freed objects, so real cells were
    # skipped at random (lost Department / Unit / responsibilities).
    seen_cells = []

    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            if any(tc is seen for seen in seen_cells):
                continue
            seen_cells.append(tc)

            paras = [_clean_text(p.text) for p in cell.paragraphs if p.text.strip()]
            if not paras:
                continue

            if _is_resp_header(paras[0]):
                resp_cell = cell
                continue

            for p in paras:
                if ":" in p:
                    label, _, val = p.partition(":")
                    field_paras[label.strip().lower()] = _clean_text(val)

    def get(labels):
        for label in labels:
            if label in field_paras and field_paras[label]:
                return field_paras[label]
        return ""

    job_title = get(["job title"])
    if not job_title:
        return None

    responsibilities: List[str] = []
    if resp_cell is not None:
        capturing = False
        for para in resp_cell.paragraphs:
            text = _clean_text(para.text)
            if not text:
                continue
            if _is_resp_header(text):
                capturing = True
                continue
            if not capturing:
                continue
            if _is_stop(text):
                break
            clean = re.sub(r"^[\-\*\•\s]+", "", text).rstrip(";").strip()
            if len(clean) > 8:
                responsibilities.append(clean)

    return {
        "source": "JD",
        "division": get(["division"]),
        "supervisor": get(["supervises"]),
        "unit": get(["unit"]),
        "department": get(["department"]),
        "reports_to": get(["reports directly to", "reports directly to:"]),
        "job_title": job_title,
        "job_grade": _parse_grade(get(["job grade"])),
        "job_category": get(["job category"]),
        "job_objective": get(["job objective"]),
        "responsibilities": responsibilities,
        "num_responsibilities": len(responsibilities),
    }


def _read_docx(path: Path) -> List[Dict]:
    from docx import Document
    records = []
    for table in Document(path).tables:
        rec = _parse_docx_table(table)
        if rec:
            records.append(rec)
    return records


# ─────────────────────────────────────────────────────────────
# PDF reader - one JD per "Job Details /Profile/" text block
# ─────────────────────────────────────────────────────────────

_BLOCK_RE = re.compile(r"(?:I\.?\s*)?Job Details\s*/Profile/\s*:", re.IGNORECASE)

_STOP_LABELS = [
    "Organizational Relationships", "Job Code", "Reports Directly To",
    "Reports Directly to", "Division", "Supervises", "Department",
    "Type of Employment", "Job Objective", "Job Grade", "Unit",
    "Job Category", "II.", "III.", "IV.", "V.", "VI.", "VII.",
]
_STOP_PAT = "|".join(re.escape(x) for x in _STOP_LABELS)
_BULLET = ""


def _split_blocks(text: str) -> List[str]:
    positions = [m.start() for m in _BLOCK_RE.finditer(text)]
    if not positions:
        return []
    return [
        text[positions[i]: positions[i + 1] if i + 1 < len(positions) else len(text)]
        for i in range(len(positions))
    ]


def _extract_job_title(block: str) -> Optional[str]:
    # The PDF's two-column layout wraps the title onto following lines: take everything up to
    # the next field boundary, drop blank lines, join and clean.
    m = re.search(
        r"Job\s+Title\s*[:\?：]\s*(.*?)(?=Organizational\s+Relationships|Job\s+Code)",
        block, re.IGNORECASE | re.DOTALL,
    )
    if not m:
        return None
    lines = [ln.strip() for ln in m.group(1).splitlines() if ln.strip()]
    if not lines:
        return None
    title = _clean_text(" ".join(lines))
    title = re.sub(r"\s*\d+\.\d+.*$", "", title).strip()
    return title or None


def _extract_field(text: str, label: str) -> Optional[str]:
    pattern = (
        rf"{re.escape(label)}\s*[:\?：]\s*"
        rf"(.*?)"
        rf"(?=\s*(?:{_STOP_PAT})\s*[:：]?|$)"
    )
    m = re.search(pattern, text, re.IGNORECASE | re.DOTALL)
    return _clean_text(m.group(1)) if m else None


def _extract_reports_to(block: str) -> str:
    m = re.search(
        r"Reports\s+Directly\s+[Tt]o\s*[:\?：]\s*(.+?)(?=Division|Department|Job\s+Grade|$)",
        block, re.IGNORECASE | re.DOTALL,
    )
    return _clean_text(m.group(1)) if m else ""


def _extract_objective(block: str) -> str:
    m = re.search(
        r"II\.?\s*Job\s+Objective\s*[:：]\s*(.+?)(?=III\.?|$)",
        block, re.IGNORECASE | re.DOTALL,
    )
    return _clean_text(m.group(1)) if m else ""


def _extract_responsibilities(block: str) -> List[str]:
    m_start = re.search(
        r"III\.\s*Key Job Duties and Responsibilities[^:]*:.*?(?=\n)", block, re.IGNORECASE)
    if not m_start:
        m_start = re.search(
            r"Key Job Duties and Responsibilities[^:]*:.*?(?=\n)", block, re.IGNORECASE)
    if not m_start:
        return []

    slice_start = m_start.end()
    m_end = re.search(r"\bIV\.?\s*Key\s+[Pp]erformance", block[slice_start:], re.IGNORECASE)
    if not m_end:
        m_end = re.search(r"\bIV\b", block[slice_start:], re.IGNORECASE)
    slice_end = slice_start + m_end.start() if m_end else len(block)

    responsibilities: List[str] = []
    for raw in block[slice_start:slice_end].split(_BULLET):
        clean = _clean_text(raw)
        if len(clean) < 8:
            continue
        if any(kw.lower() in clean.lower() for kw in RESP_STOP_KEYWORDS):
            break
        responsibilities.append(clean)
    return responsibilities


def _parse_pdf_block(block: str) -> Optional[Dict]:
    job_title = _extract_job_title(block)
    if not job_title:
        return None
    responsibilities = _extract_responsibilities(block)
    dept = _extract_field(block, "Department") or ""
    return {
        "source": "JD",
        "division": _extract_field(block, "Division") or "",
        "supervisor": _extract_field(block, "Supervises"),
        "unit": _extract_field(block, "Unit") or "",
        "department": dept if len(dept) > 3 else "",
        "reports_to": _extract_reports_to(block),
        "job_title": job_title,
        "job_grade": _parse_grade(_extract_field(block, "Job Grade")),
        "job_category": _extract_field(block, "Job Category") or "",
        "job_objective": _extract_objective(block),
        "responsibilities": responsibilities,
        "num_responsibilities": len(responsibilities),
    }


def _read_pdf(path: Path) -> List[Dict]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    records = []
    for block in _split_blocks(text):
        rec = _parse_pdf_block(block)
        if rec:
            records.append(rec)
    return records


# ─────────────────────────────────────────────────────────────
# The loader
# ─────────────────────────────────────────────────────────────

READERS = {".docx": _read_docx, ".doc": _read_docx, ".pdf": _read_pdf}


class JDLoader:
    """Parse a Job Description file (Word or PDF) of ANY division."""

    def __init__(self, file_path: Path):
        self.file_path = Path(file_path)
        self._df: Optional[pd.DataFrame] = None

    def _fill_blank_divisions(self, records: List[Dict]) -> None:
        """A JD without a Division cell takes the file's division - but only when every JD that
        does name one agrees on the same (canonical) division."""
        named = [r["division"] for r in records if r["division"]]
        if not named or len(named) == len(records):
            return
        canon = {registry.canon_division(d) or d for d in named}
        if len(canon) == 1:
            majority = Counter(named).most_common(1)[0][0]
            for r in records:
                if not r["division"]:
                    r["division"] = majority

    def load(self) -> pd.DataFrame:
        ext = self.file_path.suffix.lower()
        reader = READERS.get(ext)
        if reader is None:
            raise ValueError(f"Unsupported JD file type '{ext}'")
        print(f"\n📂 Parsing {self.file_path.name} ...")

        records = reader(self.file_path)
        self._fill_blank_divisions(records)
        if INFER_DEPARTMENT_FROM_UNIT:
            for r in records:
                if not r["department"]:
                    r["department"] = _infer_department(r["unit"], r["reports_to"], r["division"])

        for r in records:
            grade = str(r["job_grade"]) if r["job_grade"] is not None else "?"
            print(f"✓ [{grade:>2}] {r['job_title'][:45]:<45} | Unit: {(r['unit'] or '—')[:35]}")

        self._df = pd.DataFrame(records)
        total = int(self._df["num_responsibilities"].sum()) if len(self._df) else 0
        print(f"\n✅ Parsed {len(self._df)} JDs | {total} total responsibilities")
        return self._df

    def get_dataframe(self) -> pd.DataFrame:
        if self._df is None:
            self.load()
        return self._df
